import hashlib
import hmac
import json
import math
import os
import re
import secrets
import time
from io import BytesIO

from concurrent.futures import (
    ThreadPoolExecutor,
    as_completed,
)

from pathlib import Path

import pandas as pd
import matplotlib.pyplot as plt
import streamlit as st

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from pipeline import FeedbackValidator, AssignmentContext

from pipeline.core import (
    BATCH_SIZE,
    CATEGORIES,
    DEFAULT_CLASSIFICATION_PROMPT,
    DEFAULT_MODEL,
    DEFAULT_MODEL_LABEL,
    MAX_WORKERS,
    build_overall_analysis,
    calculate_model_cost,
    category_counts,
    classification_prompt_hash,
    estimate_processing_time,
    extract_pdf_text,
    parse_feedback_text,
    parse_memo_sections_from_text,
    process_one_feedback,
    results_to_dataframe,
    rubric_criterion_counts,
    validate_feedback_text,
)


# ============================================================
# PAGE CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="EduCodeInsight",
    page_icon=None,
    layout="wide",
    initial_sidebar_state="collapsed",
)


# ============================================================
# VISUAL THEME
# ============================================================

st.markdown(
    """
    <style>

    .stApp {
        background-color: #f4f7fb;
    }

    .block-container {
        max-width: 1280px;
        padding-top: 1.4rem;
        padding-bottom: 3rem;
    }

    h1, h2, h3, h4, h5, h6,
    [data-testid="stMarkdownContainer"] h1,
    [data-testid="stMarkdownContainer"] h2,
    [data-testid="stMarkdownContainer"] h3,
    [data-testid="stMarkdownContainer"] h4,
    [data-testid="stMarkdownContainer"] h5,
    [data-testid="stMarkdownContainer"] h6 {
        color: #000000 !important;
    }

    [data-testid="stMarkdownContainer"] p,
    [data-testid="stCaptionContainer"] {
        color: #1f2937 !important;
    }

    label,
    [data-testid="stWidgetLabel"],
    [data-testid="stWidgetLabel"] p {
        color: #000000 !important;
        font-weight: 600 !important;
        opacity: 1 !important;
    }

    .stTextInput input,
    .stTextArea textarea {
        background-color: #ffffff !important;
        color: #111827 !important;
        border-radius: 8px !important;
    }

    [data-testid="stForm"],
    [data-testid="stFileUploader"],
    [data-testid="stMetric"] {
        background-color: #ffffff;
        border: 1px solid #dbe3ec;
        border-radius: 12px;
    }

    [data-testid="stForm"] {
        padding: 22px;
        box-shadow: 0 5px 18px rgba(16, 42, 67, 0.06);
    }

    [data-testid="stMetric"] {
        padding: 14px;
    }

    [data-testid="stMetricLabel"],
    [data-testid="stMetricLabel"] p,
    [data-testid="stMetricValue"] {
        color: #000000 !important;
    }

    .stButton > button {
        border-radius: 8px;
        min-height: 2.6rem;
        font-weight: 600;
    }

    /* -------------------------------------------------------
       TOP NAVIGATION
       ------------------------------------------------------- */

    .st-key-top_navigation .stButton > button {
        background-color: #e9eef5 !important;
        color: #1f2937 !important;
        border: 1px solid #cbd5e1 !important;
    }

    .st-key-top_navigation .stButton > button:hover {
        background-color: #ffffff !important;
        color: #000000 !important;
        border: 1px solid #94a3b8 !important;
    }

    .st-key-top_navigation .stButton > button[kind="primary"] {
        background-color: #ffffff !important;
        color: #000000 !important;
        border: 2px solid #1f2937 !important;
        box-shadow: 0 2px 8px rgba(15, 23, 42, 0.08) !important;
    }

    .st-key-top_navigation .stButton > button[kind="primary"]:hover {
        background-color: #ffffff !important;
        color: #000000 !important;
        border: 2px solid #1f2937 !important;
    }

    footer {
        visibility: hidden;
    }

    
    /* -------------------------------------------------------
       BUTTON TEXT / CONTRAST
       ------------------------------------------------------- */

    .stButton > button,
    .stDownloadButton > button {
        color: #ffffff !important;
        background-color: #1f2937 !important;
        border: 1px solid #1f2937 !important;
        font-weight: 700 !important;
    }

    .stButton > button:hover,
    .stDownloadButton > button:hover {
        color: #ffffff !important;
        background-color: #111827 !important;
        border-color: #111827 !important;
    }

    .stButton > button[kind="primary"] {
        color: #ffffff !important;
        background-color: #ff4b4b !important;
        border-color: #ff4b4b !important;
    }

    .stButton > button[kind="primary"]:hover {
        color: #ffffff !important;
        background-color: #e53e3e !important;
        border-color: #e53e3e !important;
    }

    /* Active page stays white with dark text. */
    .st-key-top_navigation .stButton > button[kind="primary"] {
        background-color: #ffffff !important;
        color: #000000 !important;
        border: 2px solid #1f2937 !important;
        box-shadow: 0 2px 8px rgba(15, 23, 42, 0.08) !important;
    }

    /* Inactive top-navigation buttons stay dark with white text. */
    .st-key-top_navigation .stButton > button[kind="secondary"] {
        background-color: #1f2937 !important;
        color: #ffffff !important;
        border: 1px solid #1f2937 !important;
    }


    /* -------------------------------------------------------
       DOWNLOAD BUTTONS - ALWAYS READABLE
       ------------------------------------------------------- */

    .stDownloadButton > button,
    .stDownloadButton > button p,
    .stDownloadButton > button div,
    .stDownloadButton > button span {
        color: #000000 !important;
        background-color: #ffffff !important;
        border-color: #1f2937 !important;
        opacity: 1 !important;
        -webkit-text-fill-color: #000000 !important;
    }

    .stDownloadButton > button {
        border: 2px solid #1f2937 !important;
        font-weight: 700 !important;
    }

    .stDownloadButton > button:hover,
    .stDownloadButton > button:hover p,
    .stDownloadButton > button:hover div,
    .stDownloadButton > button:hover span {
        color: #000000 !important;
        background-color: #f8fafc !important;
        border-color: #111827 !important;
        -webkit-text-fill-color: #000000 !important;
    }

    .stDownloadButton > button:disabled,
    .stDownloadButton > button:disabled p,
    .stDownloadButton > button:disabled div,
    .stDownloadButton > button:disabled span {
        color: #6b7280 !important;
        background-color: #e5e7eb !important;
        border-color: #cbd5e1 !important;
        opacity: 1 !important;
        -webkit-text-fill-color: #6b7280 !important;
    }


    /* -------------------------------------------------------
       UNIFIED BUTTONS
       ------------------------------------------------------- */

    .stButton > button,
    .stButton > button p,
    .stButton > button div,
    .stButton > button span,
    .stDownloadButton > button,
    .stDownloadButton > button p,
    .stDownloadButton > button div,
    .stDownloadButton > button span {
        background-color: #e5e7eb !important;
        color: #111827 !important;
        border: 1px solid #cbd5e1 !important;
        font-weight: 700 !important;
        opacity: 1 !important;
        -webkit-text-fill-color: #111827 !important;
    }

    .stButton > button:hover,
    .stButton > button:hover p,
    .stButton > button:hover div,
    .stButton > button:hover span,
    .stDownloadButton > button:hover,
    .stDownloadButton > button:hover p,
    .stDownloadButton > button:hover div,
    .stDownloadButton > button:hover span {
        background-color: #d1d5db !important;
        color: #000000 !important;
        border-color: #94a3b8 !important;
        -webkit-text-fill-color: #000000 !important;
    }

    .stButton > button:disabled,
    .stDownloadButton > button:disabled {
        background-color: #f3f4f6 !important;
        color: #9ca3af !important;
        border-color: #e5e7eb !important;
        opacity: 1 !important;
    }

    .stButton > button:disabled p,
    .stButton > button:disabled div,
    .stButton > button:disabled span,
    .stDownloadButton > button:disabled p,
    .stDownloadButton > button:disabled div,
    .stDownloadButton > button:disabled span {
        color: #9ca3af !important;
        -webkit-text-fill-color: #9ca3af !important;
    }

    /* Active top navigation page */
    .st-key-top_navigation .stButton > button[kind="primary"],
    .st-key-top_navigation .stButton > button[kind="primary"] p,
    .st-key-top_navigation .stButton > button[kind="primary"] div,
    .st-key-top_navigation .stButton > button[kind="primary"] span {
        background-color: #ffffff !important;
        color: #000000 !important;
        border: 2px solid #111827 !important;
        box-shadow: 0 2px 8px rgba(15, 23, 42, 0.08) !important;
        -webkit-text-fill-color: #000000 !important;
    }

    /* -------------------------------------------------------
       TABS - LIGHT/NATIVE STYLE WITH ACTIVE HIGHLIGHT
       ------------------------------------------------------- */

    div[data-baseweb="tab-list"] {
        gap: 0.35rem !important;
        background: transparent !important;
    }

    button[data-baseweb="tab"] {
        background-color: transparent !important;
        color: #374151 !important;
        border-radius: 0 !important;
        border: none !important;
        font-weight: 600 !important;
        padding-left: 0.75rem !important;
        padding-right: 0.75rem !important;
    }

    button[data-baseweb="tab"] p,
    button[data-baseweb="tab"] span,
    button[data-baseweb="tab"] div {
        color: #374151 !important;
        -webkit-text-fill-color: #374151 !important;
    }

    button[data-baseweb="tab"][aria-selected="true"] {
        background-color: #ffffff !important;
        color: #000000 !important;
        border-bottom: 3px solid #111827 !important;
        font-weight: 800 !important;
    }

    button[data-baseweb="tab"][aria-selected="true"] p,
    button[data-baseweb="tab"][aria-selected="true"] span,
    button[data-baseweb="tab"][aria-selected="true"] div {
        color: #000000 !important;
        -webkit-text-fill-color: #000000 !important;
    }


    /* =======================================================
       FINAL BUTTON / TAB CONTRAST OVERRIDE
       ======================================================= */

    /* Every normal and download button: dark background, white text */
    div[data-testid="stButton"] button,
    div[data-testid="stDownloadButton"] button,
    .stButton > button,
    .stDownloadButton > button {
        background: #243041 !important;
        background-color: #243041 !important;
        color: #ffffff !important;
        border: 1px solid #243041 !important;
        font-weight: 700 !important;
        opacity: 1 !important;
        box-shadow: none !important;
        -webkit-text-fill-color: #ffffff !important;
    }

    /* Do NOT give Streamlit's internal label elements their own background/border. */
    div[data-testid="stButton"] button *,
    div[data-testid="stDownloadButton"] button *,
    .stButton > button *,
    .stDownloadButton > button * {
        background: transparent !important;
        background-color: transparent !important;
        color: #ffffff !important;
        border: none !important;
        opacity: 1 !important;
        -webkit-text-fill-color: #ffffff !important;
    }

    div[data-testid="stButton"] button:hover,
    div[data-testid="stDownloadButton"] button:hover,
    .stButton > button:hover,
    .stDownloadButton > button:hover {
        background: #334155 !important;
        background-color: #334155 !important;
        color: #ffffff !important;
        border-color: #334155 !important;
    }

    div[data-testid="stButton"] button:hover *,
    div[data-testid="stDownloadButton"] button:hover *,
    .stButton > button:hover *,
    .stDownloadButton > button:hover * {
        color: #ffffff !important;
        -webkit-text-fill-color: #ffffff !important;
    }

    /* Disabled controls stay readable but clearly disabled */
    div[data-testid="stButton"] button:disabled,
    div[data-testid="stDownloadButton"] button:disabled,
    .stButton > button:disabled,
    .stDownloadButton > button:disabled {
        background: #e5e7eb !important;
        background-color: #e5e7eb !important;
        color: #6b7280 !important;
        border-color: #cbd5e1 !important;
        opacity: 1 !important;
        -webkit-text-fill-color: #6b7280 !important;
    }

    div[data-testid="stButton"] button:disabled *,
    div[data-testid="stDownloadButton"] button:disabled *,
    .stButton > button:disabled *,
    .stDownloadButton > button:disabled * {
        color: #6b7280 !important;
        -webkit-text-fill-color: #6b7280 !important;
    }

    /* Current top-navigation page: white with black text */
    .st-key-top_navigation div[data-testid="stButton"] button[kind="primary"],
    .st-key-top_navigation .stButton > button[kind="primary"] {
        background: #ffffff !important;
        background-color: #ffffff !important;
        color: #000000 !important;
        border: 2px solid #111827 !important;
        box-shadow: 0 2px 8px rgba(15, 23, 42, 0.08) !important;
        -webkit-text-fill-color: #000000 !important;
    }

    .st-key-top_navigation div[data-testid="stButton"] button[kind="primary"] *,
    .st-key-top_navigation .stButton > button[kind="primary"] * {
        background: transparent !important;
        background-color: transparent !important;
        color: #000000 !important;
        border: none !important;
        -webkit-text-fill-color: #000000 !important;
    }

    /* Tabs: original light appearance; selected tab is clearly highlighted */
    div[data-baseweb="tab-list"] {
        background: transparent !important;
    }

    button[data-baseweb="tab"] {
        background: transparent !important;
        background-color: transparent !important;
        color: #374151 !important;
        border: none !important;
        font-weight: 600 !important;
    }

    button[data-baseweb="tab"] * {
        background: transparent !important;
        background-color: transparent !important;
        color: #374151 !important;
        border: none !important;
        -webkit-text-fill-color: #374151 !important;
    }

    button[data-baseweb="tab"][aria-selected="true"] {
        background: #ffffff !important;
        background-color: #ffffff !important;
        color: #000000 !important;
        border-bottom: 3px solid #243041 !important;
        font-weight: 800 !important;
    }

    button[data-baseweb="tab"][aria-selected="true"] * {
        color: #000000 !important;
        -webkit-text-fill-color: #000000 !important;
    }

</style>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# STORAGE
# ============================================================

DATA_ROOT = Path(
    "data"
)

USERS_FILE = (
    DATA_ROOT
    /
    "users.json"
)

USER_FILES_ROOT = (
    DATA_ROOT
    /
    "users"
)

DATA_ROOT.mkdir(
    exist_ok=True
)

USER_FILES_ROOT.mkdir(
    parents=True,
    exist_ok=True,
)

if not USERS_FILE.exists():
    USERS_FILE.write_text(
        "{}",
        encoding="utf-8",
    )


# ============================================================
# GENERAL HELPERS
# ============================================================

def get_secret(
    name,
    default=None,
):
    try:
        return st.secrets.get(
            name,
            default,
        )
    except Exception:
        return default


def get_together_api_key():
    return (
        get_secret(
            "TOGETHER_API_KEY"
        )
        or
        os.environ.get(
            "TOGETHER_API_KEY"
        )
    )


def slugify(
    value: str,
) -> str:
    value = re.sub(
        r"[^A-Za-z0-9_-]+",
        "_",
        value.strip(),
    )

    value = value.strip(
        "_"
    )

    return (
        value
        or
        "assignment"
    )


def normalize_email(
    email: str,
) -> str:
    return (
        email
        .strip()
        .lower()
    )


def valid_email(
    email: str,
) -> bool:
    return bool(
        re.fullmatch(
            r"[^\s@]+@[^\s@]+\.[^\s@]+",
            normalize_email(
                email
            ),
        )
    )


def sha256_text(
    value: str,
) -> str:
    return hashlib.sha256(
        value.encode(
            "utf-8"
        )
    ).hexdigest()


def make_cache_key(
    raw_text: str,
    memo_bytes: bytes,
    spec_bytes: bytes,
    prompt: str,
) -> str:
    digest = hashlib.sha256()

    digest.update(
        raw_text.encode(
            "utf-8"
        )
    )

    digest.update(
        memo_bytes
    )

    digest.update(
        spec_bytes
    )

    digest.update(
        DEFAULT_MODEL.encode(
            "utf-8"
        )
    )

    digest.update(
        classification_prompt_hash(
            prompt
        ).encode(
            "utf-8"
        )
    )

    digest.update(
        str(
            BATCH_SIZE
        ).encode(
            "utf-8"
        )
    )

    return digest.hexdigest()


# ============================================================
# ACCOUNT HELPERS
# ============================================================

def load_users() -> dict:
    try:
        data = json.loads(
            USERS_FILE.read_text(
                encoding="utf-8"
            )
        )

        if isinstance(
            data,
            dict,
        ):
            return data

    except Exception:
        pass

    return {}


def save_users(
    users: dict,
) -> None:
    temp_path = (
        USERS_FILE
        .with_suffix(
            ".tmp"
        )
    )

    temp_path.write_text(
        json.dumps(
            users,
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    os.replace(
        temp_path,
        USERS_FILE,
    )


def hash_password(
    password: str,
    salt: bytes,
) -> str:
    return hashlib.pbkdf2_hmac(
        "sha256",
        password.encode(
            "utf-8"
        ),
        salt,
        310_000,
    ).hex()


def create_lecturer_account(
    full_name: str,
    email: str,
    password: str,
):
    full_name = (
        full_name.strip()
    )

    email = (
        normalize_email(
            email
        )
    )

    if len(
        full_name
    ) < 2:
        return (
            False,
            "Enter your full name.",
        )

    if not valid_email(
        email
    ):
        return (
            False,
            "Enter a valid email address.",
        )

    if len(
        password
    ) < 8:
        return (
            False,
            "Password must contain at least 8 characters.",
        )

    users = (
        load_users()
    )

    if email in users:
        return (
            False,
            "An account with this email already exists.",
        )

    salt = (
        secrets.token_bytes(
            16
        )
    )

    user_id = (
        hashlib.sha256(
            email.encode(
                "utf-8"
            )
        )
        .hexdigest()[:16]
    )

    users[
        email
    ] = {
        "user_id":
            user_id,

        "full_name":
            full_name,

        "email":
            email,

        "password_salt":
            salt.hex(),

        "password_hash":
            hash_password(
                password,
                salt,
            ),

        "created_at":
            time.strftime(
                "%Y-%m-%d %H:%M:%S"
            ),
    }

    save_users(
        users
    )

    root = (
        USER_FILES_ROOT
        /
        user_id
    )

    (
        root
        /
        "assignments"
    ).mkdir(
        parents=True,
        exist_ok=True,
    )

    (
        root
        /
        "reports"
    ).mkdir(
        parents=True,
        exist_ok=True,
    )

    return (
        True,
        "Account created successfully. You can now log in.",
    )


def authenticate_lecturer(
    email: str,
    password: str,
):
    users = (
        load_users()
    )

    user = users.get(
        normalize_email(
            email
        )
    )

    if not user:
        return None

    try:
        salt = bytes.fromhex(
            user[
                "password_salt"
            ]
        )

        supplied = (
            hash_password(
                password,
                salt,
            )
        )

    except Exception:
        return None

    if hmac.compare_digest(
        user[
            "password_hash"
        ],
        supplied,
    ):
        return user

    return None


# ============================================================
# USER STORAGE
# ============================================================

def current_user_root():
    user_id = (
        st.session_state.get(
            "user_id"
        )
    )

    if not user_id:
        raise RuntimeError(
            "No lecturer is currently authenticated."
        )

    root = (
        USER_FILES_ROOT
        /
        str(
            user_id
        )
    )

    root.mkdir(
        parents=True,
        exist_ok=True,
    )

    return root


def assignments_root():
    folder = (
        current_user_root()
        /
        "assignments"
    )

    folder.mkdir(
        parents=True,
        exist_ok=True,
    )

    return folder


def reports_root():
    folder = (
        current_user_root()
        /
        "reports"
    )

    folder.mkdir(
        parents=True,
        exist_ok=True,
    )

    return folder


def assignment_folder(
    assignment_name: str,
):
    folder = (
        assignments_root()
        /
        slugify(
            assignment_name
        )
    )

    folder.mkdir(
        parents=True,
        exist_ok=True,
    )

    return folder


def assignment_meta_path(
    assignment_name: str,
):
    return (
        assignment_folder(
            assignment_name
        )
        /
        "_assignment.json"
    )


def list_stored_assignments():
    assignments = []

    for folder in (
        assignments_root()
        .iterdir()
    ):
        if not folder.is_dir():
            continue

        meta_path = (
            folder
            /
            "_assignment.json"
        )

        if not meta_path.exists():
            continue

        try:
            meta = json.loads(
                meta_path.read_text(
                    encoding="utf-8"
                )
            )

            assignments.append(
                (
                    meta.get(
                        "assignment_name",
                        folder.name,
                    ),
                    folder.name,
                )
            )

        except Exception:
            continue

    assignments.sort(
        key=lambda item:
            item[0].lower()
    )

    return assignments


def load_assignment_by_slug(
    assignment_slug: str,
):
    folder = (
        assignments_root()
        /
        assignment_slug
    )

    meta = {}

    meta_path = (
        folder
        /
        "_assignment.json"
    )

    if meta_path.exists():
        try:
            meta = json.loads(
                meta_path.read_text(
                    encoding="utf-8"
                )
            )
        except Exception:
            meta = {}

    results = []

    for result_path in sorted(
        folder.glob(
            "*.json"
        )
    ):
        if result_path.name == "_assignment.json":
            continue

        try:
            result = json.loads(
                result_path.read_text(
                    encoding="utf-8"
                )
            )

            result.pop(
                "_cache_key",
                None,
            )

            results.append(
                result
            )

        except Exception:
            continue

    return (
        results,
        meta,
    )


def load_all_assignment_runs():
    runs = {}

    for (
        assignment_name,
        assignment_slug,
    ) in list_stored_assignments():

        results, meta = (
            load_assignment_by_slug(
                assignment_slug
            )
        )

        if not results:
            continue

        runs[
            assignment_name
        ] = {
            **meta,
            "results":
                results,

            "processed_count":
                meta.get(
                    "processed_count",
                    len(
                        results
                    ),
                ),
        }

    return runs


# ============================================================
# ESTIMATION HELPERS
# ============================================================

def estimate_pre_run_usage(
    valid_feedback,
    memo_text: str,
    specification_text: str,
    classification_prompt: str,
):
    """
    Provide a transparent pre-run token/cost estimate.

    This estimate uses character counts and the number of
    failed/partial rubric items. Actual provider tokenisation,
    model reasoning output, retries and latency may differ.
    """

    estimated_requests = 0
    estimated_failed_items = 0
    estimated_input_chars = 0

    prompt_chars = len(
        classification_prompt
    )

    memo_chars = min(
        len(
            memo_text
        ),
        7000,
    )

    spec_chars = min(
        len(
            specification_text
        ),
        7000,
    )

    for (
        filename,
        raw_text,
    ) in valid_feedback:

        try:
            parsed = parse_feedback_text(
                raw_text,
                filename,
            )
        except Exception:
            continue

        for question in parsed.questions:

            failed_items = [
                item
                for item
                in question.items
                if not item.met
            ]

            if not failed_items:
                continue

            estimated_failed_items += len(
                failed_items
            )

            requests = math.ceil(
                len(
                    failed_items
                )
                /
                BATCH_SIZE
            )

            estimated_requests += (
                requests
            )

            item_chars = sum(
                len(
                    item.text
                    or
                    ""
                )
                +
                len(
                    item.reason
                    or
                    ""
                )
                for item
                in failed_items
            )

            estimated_input_chars += (
                requests
                *
                (
                    prompt_chars
                    +
                    memo_chars
                    +
                    spec_chars
                    +
                    len(
                        question.title
                        or
                        ""
                    )
                )
                +
                item_chars
            )

    estimated_input_tokens = math.ceil(
        estimated_input_chars
        /
        4
    )

    estimated_output_tokens = max(
        estimated_failed_items
        *
        90,
        estimated_requests
        *
        120,
    )

    estimated_cost = (
        calculate_model_cost(
            DEFAULT_MODEL,
            estimated_input_tokens,
            estimated_output_tokens,
        )
    )

    return {
        "failed_partial_items":
            estimated_failed_items,

        "api_requests":
            estimated_requests,

        "input_tokens":
            estimated_input_tokens,

        "output_tokens":
            estimated_output_tokens,

        "total_tokens":
            (
                estimated_input_tokens
                +
                estimated_output_tokens
            ),

        "cost_usd":
            estimated_cost,
    }


# ============================================================
# REPORT HELPERS
# ============================================================

def _format_docx_document(
    document: Document,
):
    """
    Apply restrained formatting to EduCodeInsight reports.
    """

    styles = document.styles

    styles[
        "Normal"
    ].font.name = "Aptos"

    styles[
        "Normal"
    ].font.size = Pt(
        10
    )

    styles[
        "Title"
    ].font.name = "Aptos Display"

    styles[
        "Title"
    ].font.size = Pt(
        22
    )

    styles[
        "Heading 1"
    ].font.name = "Aptos Display"

    styles[
        "Heading 1"
    ].font.size = Pt(
        15
    )

    styles[
        "Heading 2"
    ].font.name = "Aptos Display"

    styles[
        "Heading 2"
    ].font.size = Pt(
        12
    )


def _add_key_value_table(
    document: Document,
    rows,
):
    table = document.add_table(
        rows=1,
        cols=2,
    )

    table.style = "Table Grid"

    table.rows[
        0
    ].cells[
        0
    ].text = "Measure"

    table.rows[
        0
    ].cells[
        1
    ].text = "Value"

    for label, value in rows:

        cells = (
            table.add_row().cells
        )

        cells[
            0
        ].text = str(
            label
        )

        cells[
            1
        ].text = str(
            value
        )

    document.add_paragraph()


def _shorten_text(
    value,
    maximum: int = 180,
) -> str:
    value = (
        ""
        if value is None
        else str(
            value
        )
    )

    value = re.sub(
        r"\s+",
        " ",
        value,
    ).strip()

    if len(
        value
    ) <= maximum:
        return value

    return (
        value[
            : maximum - 1
        ].rstrip()
        +
        "…"
    )


def _figure_bytes(
    fig,
) -> bytes:
    buffer = BytesIO()

    fig.savefig(
        buffer,
        format="png",
        dpi=180,
        bbox_inches="tight",
    )

    plt.close(
        fig
    )

    buffer.seek(
        0
    )

    return buffer.getvalue()


def _category_chart_bytes(
    counts,
) -> bytes:
    series = (
        counts
        .reindex(
            CATEGORIES,
            fill_value=0,
        )
    )

    fig, ax = plt.subplots(
        figsize=(
            7.2,
            4.2,
        )
    )

    ax.bar(
        series.index,
        series.values,
    )

    ax.set_title(
        "Error category distribution"
    )

    ax.set_ylabel(
        "Failed / partial criteria"
    )

    ax.tick_params(
        axis="x",
        rotation=20,
    )

    fig.tight_layout()

    return _figure_bytes(
        fig
    )


def _top_patterns_chart_bytes(
    patterns,
    top_n: int = 8,
) -> bytes:
    top = (
        patterns
        .head(
            top_n
        )
        .copy()
    )

    if top.empty:
        return b""

    top[
        "label"
    ] = top[
        "rubric_text"
    ].apply(
        lambda value:
            _shorten_text(
                value,
                62,
            )
    )

    top = (
        top
        .sort_values(
            "count",
            ascending=True,
        )
    )

    fig, ax = plt.subplots(
        figsize=(
            8.0,
            5.0,
        )
    )

    ax.barh(
        top[
            "label"
        ],
        top[
            "count"
        ],
    )

    ax.set_title(
        "Most common recurring errors"
    )

    ax.set_xlabel(
        "Students affected"
    )

    fig.tight_layout()

    return _figure_bytes(
        fig
    )


def _question_chart_bytes(
    df,
) -> bytes:
    if (
        df.empty
        or
        "question_number"
        not in df.columns
    ):
        return b""

    question_counts = (
        df
        .groupby(
            "question_number",
            dropna=False,
        )
        .size()
        .reset_index(
            name="count"
        )
    )

    if question_counts.empty:
        return b""

    question_counts[
        "question"
    ] = question_counts[
        "question_number"
    ].apply(
        lambda value:
            (
                "Question "
                +
                str(
                    value
                )
            )
    )

    fig, ax = plt.subplots(
        figsize=(
            6.8,
            3.8,
        )
    )

    ax.bar(
        question_counts[
            "question"
        ],
        question_counts[
            "count"
        ],
    )

    ax.set_title(
        "Failed / partial criteria by question"
    )

    ax.set_ylabel(
        "Count"
    )

    fig.tight_layout()

    return _figure_bytes(
        fig
    )


def _assignment_main_findings(
    results,
    df,
    counts,
    patterns,
) -> list:
    findings = []

    total_items = len(
        df
    )

    students = len(
        results
    )

    if total_items == 0:

        return [
            (
                f"{students} student submissions were analysed and "
                "no failed or partially met rubric criteria were found."
            )
        ]

    findings.append(
        (
            f"{students} student submissions were analysed, producing "
            f"{total_items} failed or partially met rubric criteria."
        )
    )

    if int(
        counts.sum()
    ) > 0:

        dominant_category = (
            counts.idxmax()
        )

        dominant_count = int(
            counts.max()
        )

        dominant_percentage = (
            dominant_count
            /
            total_items
            *
            100
        )

        findings.append(
            (
                f"The dominant error type was {dominant_category}: "
                f"{dominant_count} items ({dominant_percentage:.1f}% "
                "of all failed/partial criteria)."
            )
        )

    if not patterns.empty:

        first = (
            patterns.iloc[
                0
            ]
        )

        findings.append(
            (
                "The most widespread recurring issue was "
                f"“{_shorten_text(first['rubric_text'], 150)}”, "
                f"which affected {int(first['count'])} students."
            )
        )

        if len(
            patterns
        ) > 1:

            next_items = []

            for _, row in (
                patterns
                .iloc[
                    1:5
                ]
                .iterrows()
            ):

                next_items.append(
                    (
                        f"“{_shorten_text(row['rubric_text'], 105)}” "
                        f"({int(row['count'])} students)"
                    )
                )

            if next_items:

                findings.append(
                    (
                        "Other common areas of difficulty were "
                        +
                        "; ".join(
                            next_items
                        )
                        +
                        "."
                    )
                )

    return findings


def assignment_report_docx(
    assignment_name: str,
    results,
    meta: dict,
) -> bytes:
    """
    Build a concise visual Word report for one assignment.

    Detailed student-level rows remain available in the Streamlit
    dashboard instead of being repeated in the report.
    """

    df = (
        results_to_dataframe(
            results
        )
    )

    counts = (
        category_counts(
            df
        )
    )

    patterns = (
        rubric_criterion_counts(
            df
        )
    )

    findings = (
        _assignment_main_findings(
            results,
            df,
            counts,
            patterns,
        )
    )

    document = Document()

    _format_docx_document(
        document
    )

    title = document.add_heading(
        "EduCodeInsight Assignment Analysis Report",
        level=0,
    )

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    subtitle = (
        document.add_paragraph()
    )

    subtitle.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    subtitle.add_run(
        assignment_name
    ).bold = True


    # --------------------------------------------------------
    # EXECUTIVE SUMMARY
    # --------------------------------------------------------

    document.add_heading(
        "1. Executive Summary",
        level=1,
    )

    for finding in findings:

        paragraph = (
            document.add_paragraph(
                style="List Bullet"
            )
        )

        paragraph.add_run(
            finding
        )


    # --------------------------------------------------------
    # VISUAL OVERVIEW
    # --------------------------------------------------------

    document.add_heading(
        "2. Visual Overview",
        level=1,
    )

    category_image = (
        _category_chart_bytes(
            counts
        )
    )

    document.add_picture(
        BytesIO(
            category_image
        ),
        width=Inches(
            6.4
        ),
    )

    document.add_paragraph(
        (
            "This chart shows the distribution of the four "
            "EduCodeInsight error categories for all failed or "
            "partially met rubric criteria."
        )
    )

    question_image = (
        _question_chart_bytes(
            df
        )
    )

    if question_image:

        document.add_picture(
            BytesIO(
                question_image
            ),
            width=Inches(
                6.2
            ),
        )

        document.add_paragraph(
            (
                "This chart shows where failed or partially met "
                "criteria were concentrated across assignment questions."
            )
        )

    pattern_image = (
        _top_patterns_chart_bytes(
            patterns
        )
    )

    if pattern_image:

        document.add_picture(
            BytesIO(
                pattern_image
            ),
            width=Inches(
                6.5
            ),
        )

        document.add_paragraph(
            (
                "This chart highlights the recurring rubric criteria "
                "that affected the largest number of students."
            )
        )


    # --------------------------------------------------------
    # MAIN AREAS THAT WERE INCORRECT
    # --------------------------------------------------------

    document.add_heading(
        "3. What Students Mainly Got Wrong",
        level=1,
    )

    if patterns.empty:

        document.add_paragraph(
            "No recurring error patterns were identified."
        )

    else:

        top_patterns = (
            patterns
            .head(
                10
            )
        )

        table = document.add_table(
            rows=1,
            cols=3,
        )

        table.style = (
            "Table Grid"
        )

        headers = [
            "Common Issue",
            "Students Affected",
            "Dominant Error Type",
        ]

        for index, header in enumerate(
            headers
        ):

            table.rows[
                0
            ].cells[
                index
            ].text = header

        for _, row in (
            top_patterns.iterrows()
        ):

            cells = (
                table.add_row().cells
            )

            cells[
                0
            ].text = (
                _shorten_text(
                    row[
                        "rubric_text"
                    ],
                    220,
                )
            )

            cells[
                1
            ].text = str(
                int(
                    row[
                        "count"
                    ]
                )
            )

            cells[
                2
            ].text = str(
                row[
                    "dominant_category"
                ]
            )


    # --------------------------------------------------------
    # CATEGORY SUMMARY
    # --------------------------------------------------------

    document.add_paragraph()

    document.add_heading(
        "4. Error Category Summary",
        level=1,
    )

    category_table = (
        document.add_table(
            rows=1,
            cols=3,
        )
    )

    category_table.style = (
        "Table Grid"
    )

    headers = [
        "Error Category",
        "Count",
        "Share",
    ]

    for index, header in enumerate(
        headers
    ):

        category_table.rows[
            0
        ].cells[
            index
        ].text = header

    total = max(
        int(
            counts.sum()
        ),
        1,
    )

    for category, count in (
        counts.items()
    ):

        cells = (
            category_table
            .add_row()
            .cells
        )

        cells[
            0
        ].text = str(
            category
        )

        cells[
            1
        ].text = str(
            int(
                count
            )
        )

        cells[
            2
        ].text = (
            f"{int(count) / total * 100:.1f}%"
        )


    # --------------------------------------------------------
    # TECHNICAL DETAILS
    # --------------------------------------------------------

    document.add_paragraph()

    document.add_heading(
        "5. Analysis Details",
        level=1,
    )

    technical_rows = [
        (
            "Students processed",
            len(
                results
            ),
        ),
        (
            "Failed / partial rubric items",
            len(
                df
            ),
        ),
        (
            "Model",
            meta.get(
                "model_label",
                DEFAULT_MODEL_LABEL,
            ),
        ),
        (
            "Processing time",
            (
                f"{meta.get('elapsed_minutes', 0.0):.2f} minutes"
            ),
        ),
        (
            "Prompt modified",
            meta.get(
                "prompt_modified",
                False,
            ),
        ),
        (
            "Prompt hash",
            meta.get(
                "prompt_hash",
                "",
            ),
        ),
    ]

    if (
        meta.get(
            "dataset_total_tokens",
            0,
        )
        or
        meta.get(
            "dataset_cost_usd",
            0.0,
        )
    ):

        technical_rows.extend(
            [
                (
                    "Total tokens",
                    f"{meta.get('dataset_total_tokens', 0):,}",
                ),
                (
                    "AI cost",
                    f"${meta.get('dataset_cost_usd', 0.0):.6f}",
                ),
            ]
        )

    _add_key_value_table(
        document,
        technical_rows,
    )

    document.add_paragraph(
        (
            "Detailed student-level classifications are intentionally "
            "kept in the EduCodeInsight dashboard rather than repeated "
            "in this summary report."
        )
    )

    buffer = (
        BytesIO()
    )

    document.save(
        buffer
    )

    return buffer.getvalue()


def _overall_category_chart_bytes(
    category_distribution,
) -> bytes:
    if category_distribution.empty:
        return b""

    fig, ax = plt.subplots(
        figsize=(
            7.0,
            4.0,
        )
    )

    ax.bar(
        category_distribution[
            "category"
        ],
        category_distribution[
            "count"
        ],
    )

    ax.set_title(
        "Overall error category distribution"
    )

    ax.set_ylabel(
        "Count"
    )

    ax.tick_params(
        axis="x",
        rotation=20,
    )

    fig.tight_layout()

    return _figure_bytes(
        fig
    )


def _assignment_comparison_chart_bytes(
    assignment_summary,
) -> bytes:
    if assignment_summary.empty:
        return b""

    data = (
        assignment_summary
        .sort_values(
            "failed_partial_items",
            ascending=True,
        )
    )

    fig, ax = plt.subplots(
        figsize=(
            7.6,
            4.8,
        )
    )

    ax.barh(
        data[
            "assignment"
        ],
        data[
            "failed_partial_items"
        ],
    )

    ax.set_title(
        "Failed / partial criteria by assignment"
    )

    ax.set_xlabel(
        "Count"
    )

    fig.tight_layout()

    return _figure_bytes(
        fig
    )


def overall_report_docx(
    overall: dict,
) -> bytes:
    """
    Build a concise visual report across all stored assignments.
    """

    summary = (
        overall[
            "summary"
        ]
    )

    assignment_summary = (
        overall[
            "assignment_summary"
        ]
    )

    category_distribution = (
        overall[
            "category_distribution"
        ]
    )

    recurring_errors = (
        overall[
            "recurring_errors"
        ]
    )

    document = Document()

    _format_docx_document(
        document
    )

    title = document.add_heading(
        "EduCodeInsight Overall Analysis Report",
        level=0,
    )

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )


    document.add_heading(
        "1. Executive Summary",
        level=1,
    )

    summary_points = [
        (
            f"{summary['assignments_analysed']} assignments and "
            f"{summary['student_submissions_processed']} student "
            "submissions were included in the overall analysis."
        ),
        (
            f"The combined dataset contains "
            f"{summary['failed_partial_items']} failed or partially "
            "met rubric criteria."
        ),
    ]

    if not category_distribution.empty:

        dominant_row = (
            category_distribution
            .sort_values(
                "count",
                ascending=False,
            )
            .iloc[
                0
            ]
        )

        total_categories = max(
            int(
                category_distribution[
                    "count"
                ].sum()
            ),
            1,
        )

        summary_points.append(
            (
                f"The dominant error type across assignments was "
                f"{dominant_row['category']} with "
                f"{int(dominant_row['count'])} occurrences "
                f"({int(dominant_row['count']) / total_categories * 100:.1f}%)."
            )
        )

    if not recurring_errors.empty:

        top = (
            recurring_errors.iloc[
                0
            ]
        )

        summary_points.append(
            (
                "The most persistent recurring issue across assignments "
                f"was “{_shorten_text(top['rubric_text'], 150)}”."
            )
        )

    for point in summary_points:

        paragraph = (
            document.add_paragraph(
                style="List Bullet"
            )
        )

        paragraph.add_run(
            point
        )


    document.add_heading(
        "2. Visual Overview",
        level=1,
    )

    category_image = (
        _overall_category_chart_bytes(
            category_distribution
        )
    )

    if category_image:

        document.add_picture(
            BytesIO(
                category_image
            ),
            width=Inches(
                6.4
            ),
        )

    comparison_image = (
        _assignment_comparison_chart_bytes(
            assignment_summary
        )
    )

    if comparison_image:

        document.add_picture(
            BytesIO(
                comparison_image
            ),
            width=Inches(
                6.4
            ),
        )


    document.add_heading(
        "3. Main Recurring Problems",
        level=1,
    )

    if recurring_errors.empty:

        document.add_paragraph(
            "No recurring cross-assignment problems were identified."
        )

    else:

        table = (
            document.add_table(
                rows=1,
                cols=4,
            )
        )

        table.style = (
            "Table Grid"
        )

        headers = [
            "Common Issue",
            "Occurrences",
            "Students Affected",
            "Assignments Affected",
        ]

        for index, header in enumerate(
            headers
        ):

            table.rows[
                0
            ].cells[
                index
            ].text = header

        for _, row in (
            recurring_errors
            .head(
                12
            )
            .iterrows()
        ):

            cells = (
                table.add_row().cells
            )

            values = [
                _shorten_text(
                    row.get(
                        "rubric_text",
                        "",
                    ),
                    200,
                ),
                row.get(
                    "occurrences",
                    0,
                ),
                row.get(
                    "students_affected",
                    0,
                ),
                row.get(
                    "assignments_affected",
                    0,
                ),
            ]

            for index, value in enumerate(
                values
            ):

                cells[
                    index
                ].text = str(
                    value
                )


    document.add_paragraph()

    document.add_heading(
        "4. Assignment Comparison",
        level=1,
    )

    if assignment_summary.empty:

        document.add_paragraph(
            "No assignment comparison data is available."
        )

    else:

        table = (
            document.add_table(
                rows=1,
                cols=4,
            )
        )

        table.style = (
            "Table Grid"
        )

        headers = [
            "Assignment",
            "Students",
            "Failed / Partial Items",
            "Processing Minutes",
        ]

        for index, header in enumerate(
            headers
        ):

            table.rows[
                0
            ].cells[
                index
            ].text = header

        for _, row in (
            assignment_summary.iterrows()
        ):

            cells = (
                table.add_row().cells
            )

            values = [
                row.get(
                    "assignment",
                    "",
                ),
                row.get(
                    "students_processed",
                    0,
                ),
                row.get(
                    "failed_partial_items",
                    0,
                ),
                (
                    f"{float(row.get('processing_minutes', 0.0)):.2f}"
                ),
            ]

            for index, value in enumerate(
                values
            ):

                cells[
                    index
                ].text = str(
                    value
                )


    document.add_paragraph()

    document.add_heading(
        "5. Analysis Details",
        level=1,
    )

    detail_rows = [
        (
            "Assignments analysed",
            summary[
                "assignments_analysed"
            ],
        ),
        (
            "Student submissions",
            summary[
                "student_submissions_processed"
            ],
        ),
        (
            "Unique students",
            summary[
                "unique_students"
            ],
        ),
        (
            "Failed / partial items",
            summary[
                "failed_partial_items"
            ],
        ),
        (
            "Total processing time",
            (
                f"{summary['total_processing_minutes']:.2f} minutes"
            ),
        ),
    ]

    if (
        summary.get(
            "total_tokens",
            0,
        )
        or
        summary.get(
            "total_ai_cost_usd",
            0.0,
        )
    ):

        detail_rows.extend(
            [
                (
                    "Total tokens",
                    f"{summary['total_tokens']:,}",
                ),
                (
                    "Total AI cost",
                    f"${summary['total_ai_cost_usd']:.6f}",
                ),
            ]
        )

    _add_key_value_table(
        document,
        detail_rows,
    )

    buffer = (
        BytesIO()
    )

    document.save(
        buffer
    )

    return buffer.getvalue()


# ============================================================
# SESSION STATE
# ============================================================

defaults = {
    "authenticated":
        False,

    "user_id":
        None,

    "full_name":
        None,

    "page":
        "Home",

    "classification_prompt":
        DEFAULT_CLASSIFICATION_PROMPT.strip(),
}

for key, value in defaults.items():
    if key not in st.session_state:
        st.session_state[
            key
        ] = value


# ============================================================
# AUTHENTICATION PAGE
# ============================================================

def render_authentication():
    st.title(
        "EduCodeInsight"
    )

    st.caption(
        "Automated error analysis for student feedback"
    )

    login_tab, register_tab = st.tabs(
        [
            "Login",
            "Create Account",
        ]
    )

    with login_tab:
        with st.form(
            "login_form"
        ):
            email = st.text_input(
                "Email address"
            )

            password = st.text_input(
                "Password",
                type="password",
            )

            submitted = st.form_submit_button(
                "Login",
                use_container_width=True,
            )

        if submitted:
            user = authenticate_lecturer(
                email,
                password,
            )

            if user:
                st.session_state[
                    "authenticated"
                ] = True

                st.session_state[
                    "user_id"
                ] = user[
                    "user_id"
                ]

                st.session_state[
                    "full_name"
                ] = user[
                    "full_name"
                ]

                st.session_state[
                    "page"
                ] = "Home"

                st.rerun()

            else:
                st.error(
                    "Incorrect email address or password."
                )

    with register_tab:
        with st.form(
            "register_form"
        ):
            full_name = st.text_input(
                "Full name"
            )

            email = st.text_input(
                "Email address",
                key="register_email",
            )

            password = st.text_input(
                "Password",
                type="password",
                key="register_password",
            )

            confirm = st.text_input(
                "Confirm password",
                type="password",
            )

            submitted = st.form_submit_button(
                "Create account",
                use_container_width=True,
            )

        if submitted:
            if password != confirm:
                st.error(
                    "The passwords do not match."
                )
            else:
                success, message = (
                    create_lecturer_account(
                        full_name,
                        email,
                        password,
                    )
                )

                if success:
                    st.success(
                        message
                    )
                else:
                    st.error(
                        message
                    )


# ============================================================
# TOP NAVIGATION
# ============================================================

def render_navigation():
    st.title(
        "EduCodeInsight"
    )

    st.caption(
        f"Signed in as {st.session_state.get('full_name', 'Lecturer')}"
    )

    pages = [
        "Home",
        "Analyse",
        "Dashboard",
        "Reports",
    ]

    with st.container(
        key="top_navigation"
    ):

        columns = st.columns(
            len(
                pages
            )
            +
            1
        )

        for index, page in enumerate(
            pages
        ):

            active_page = (
                st.session_state[
                    "page"
                ]
                ==
                page
            )

            if columns[
                index
            ].button(
                page,
                use_container_width=True,
                type=(
                    "primary"
                    if active_page
                    else
                    "secondary"
                ),
                key=(
                    "nav_"
                    +
                    page.lower()
                ),
            ):

                st.session_state[
                    "page"
                ] = page

                st.rerun()

        if columns[
            -1
        ].button(
            "Logout",
            use_container_width=True,
            key="nav_logout",
        ):

            for key in [
                "authenticated",
                "user_id",
                "full_name",
            ]:

                st.session_state[
                    key
                ] = (
                    False
                    if key
                    ==
                    "authenticated"
                    else
                    None
                )

            st.session_state[
                "page"
            ] = "Home"

            st.rerun()

    st.divider()


# ============================================================
# HOME PAGE
# ============================================================

def render_home():
    st.header(
        "Home"
    )

    st.write(
        "EduCodeInsight analyses already-marked student feedback "
        "and classifies failed or partially met rubric criteria "
        "into four error categories."
    )

    c1, c2, c3 = st.columns(
        3
    )

    c1.metric(
        "Classification Model",
        DEFAULT_MODEL_LABEL,
    )

    c2.metric(
        "Stored Assignments",
        len(
            list_stored_assignments()
        ),
    )

    c3.metric(
        "Error Categories",
        len(
            CATEGORIES
        ),
    )

    st.subheader(
        "Workflow"
    )

    st.write(
        "Upload feedback files, the marking memorandum and the "
        "assignment specification on the Analyse page. Review the "
        "classification prompt and the estimated runtime/cost before "
        "starting the analysis. Results are then available in the "
        "assignment and overall Streamlit dashboards."
    )


# ============================================================
# ANALYSE PAGE
# ============================================================

def render_analyse():
    # Services receive the bytes/text supplied by the lecturer via Streamlit.
    feedback_validator = FeedbackValidator()
    assignment_context = AssignmentContext()

    st.header(
        "Analyse Assignment"
    )

    st.info(
        f"EduCodeInsight currently uses {DEFAULT_MODEL_LABEL} "
        f"({DEFAULT_MODEL}) as the single validated model."
    )

    api_key = (
        get_together_api_key()
    )

    if not api_key:
        st.error(
            "The Together AI API key has not been configured. "
            "Configure .streamlit/secrets.toml before running analysis."
        )

    assignment_name = st.text_input(
        "Assignment name",
        placeholder="Example: Practical 6",
    )

    feedback_files = st.file_uploader(
        "Student feedback files",
        type=["txt"],
        accept_multiple_files=True,
    )

    memo_file = st.file_uploader(
        "Marking memorandum",
        type=["pdf"],
    )

    spec_file = st.file_uploader(
        "Assignment specification",
        type=["pdf"],
    )

    st.subheader(
        "Classification Prompt"
    )

    st.caption(
        "The lecturer can review and edit the exact prompt that "
        "will be used for this assignment. The prompt and its hash "
        "are stored with the analysis for reproducibility."
    )

    restore_col, info_col = st.columns(
        [1, 3]
    )

    if restore_col.button(
        "Restore Default Prompt"
    ):
        st.session_state[
            "classification_prompt"
        ] = (
            DEFAULT_CLASSIFICATION_PROMPT.strip()
        )
        st.rerun()

    info_col.caption(
        "Changing the prompt changes the classification behaviour "
        "for the complete assignment run."
    )

    prompt = st.text_area(
        "Prompt",
        key="classification_prompt",
        height=390,
    )

    prompt_hash = (
        classification_prompt_hash(
            prompt
        )
    )

    prompt_modified = (
        prompt.strip()
        !=
        DEFAULT_CLASSIFICATION_PROMPT.strip()
    )

    st.caption(
        f"Prompt modified: {prompt_modified} | "
        f"SHA-256: {prompt_hash[:16]}..."
    )

    prompt_approved = st.checkbox(
        "I have reviewed this prompt and want EduCodeInsight to use it."
    )

    valid_feedback = []
    invalid_feedback = []

    memo_bytes = b""
    spec_bytes = b""
    memo_text = ""
    specification_text = ""
    memo_sections = []

    if feedback_files:
        for uploaded in feedback_files:

            raw_text = uploaded.getvalue().decode(
                "utf-8",
                errors="replace",
            )

            valid, errors = (
                feedback_validator.validate(
                    raw_text,
                    uploaded.name,
                )
            )

            if valid:
                valid_feedback.append(
                    (
                        uploaded.name,
                        raw_text,
                    )
                )
            else:
                invalid_feedback.append(
                    {
                        "file":
                            uploaded.name,
                        "errors":
                            "; ".join(
                                errors
                            ),
                    }
                )

    if memo_file is not None:
        memo_bytes = memo_file.getvalue()

        try:
            memo_text = (
                assignment_context.extract_pdf_text(
                    memo_bytes
                )
            )

            memo_sections = (
                assignment_context.parse_memo(
                    memo_text
                )
            )

        except Exception as exc:
            st.error(
                f"Could not read the memorandum: {exc}"
            )

    if spec_file is not None:
        spec_bytes = spec_file.getvalue()

        try:
            specification_text = (
                assignment_context.extract_pdf_text(
                    spec_bytes
                )
            )

        except Exception as exc:
            st.error(
                f"Could not read the specification: {exc}"
            )

    if feedback_files:
        c1, c2, c3 = st.columns(
            3
        )

        c1.metric(
            "Uploaded Feedback",
            len(
                feedback_files
            ),
        )

        c2.metric(
            "Valid Feedback",
            len(
                valid_feedback
            ),
        )

        c3.metric(
            "Invalid Feedback",
            len(
                invalid_feedback
            ),
        )

        if invalid_feedback:
            with st.expander(
                "Show invalid feedback files"
            ):
                st.dataframe(
                    pd.DataFrame(
                        invalid_feedback
                    ),
                    use_container_width=True,
                    hide_index=True,
                )

    ready_for_estimate = (
        bool(
            valid_feedback
        )
        and
        memo_file is not None
        and
        spec_file is not None
        and
        bool(
            prompt.strip()
        )
    )

    estimate = None
    time_estimate = None

    if ready_for_estimate:
        time_estimate = (
            estimate_processing_time(
                len(
                    valid_feedback
                )
            )
        )

        estimate = (
            estimate_pre_run_usage(
                valid_feedback,
                memo_text,
                specification_text,
                prompt,
            )
        )

        st.subheader(
            "Pre-run Estimate"
        )

        st.warning(
            "These values are estimates. Actual runtime and cost can "
            "change because of provider tokenisation, reasoning output, "
            "API latency, retries and rate limits."
        )

        c1, c2, c3, c4 = st.columns(
            4
        )

        c1.metric(
            "Files to Process",
            len(
                valid_feedback
            ),
        )

        c2.metric(
            "Estimated Time",
            (
                f"{time_estimate['lower_minutes']}"
                f"–{time_estimate['upper_minutes']} min"
            ),
        )

        c3.metric(
            "Estimated Tokens",
            f"{estimate['total_tokens']:,}",
        )

        c4.metric(
            "Estimated Cost",
            f"${estimate['cost_usd']:.4f}",
        )

        st.caption(
            f"Estimated API requests: {estimate['api_requests']:,} | "
            f"Estimated failed/partial criteria: "
            f"{estimate['failed_partial_items']:,} | "
            f"Workers: {MAX_WORKERS} | Batch size: {BATCH_SIZE}"
        )

    can_run = (
        bool(
            assignment_name.strip()
        )
        and
        bool(
            valid_feedback
        )
        and
        memo_file is not None
        and
        spec_file is not None
        and
        bool(
            api_key
        )
        and
        bool(
            prompt.strip()
        )
        and
        prompt_approved
    )

    if st.button(
        "Run Analysis",
        type="primary",
        use_container_width=True,
        disabled=not can_run,
    ):

        assignment_name_clean = (
            assignment_name.strip()
        )

        folder = (
            assignment_folder(
                assignment_name_clean
            )
        )

        start_time = (
            time.perf_counter()
        )

        results = []
        failures = []
        work = []

        # ====================================================
        # FORCE ALL VALID FILES INTO THE PROCESSING QUEUE
        # ====================================================

        for (
            filename,
            raw_text,
        ) in valid_feedback:

            result_path = (
                folder
                /
                (
                    slugify(
                        Path(
                            filename
                        ).stem
                    )
                    +
                    ".json"
                )
            )

            audit_key = (
                make_cache_key(
                    raw_text,
                    memo_bytes,
                    spec_bytes,
                    prompt,
                )
            )

            work.append(
                (
                    filename,
                    raw_text,
                    result_path,
                    audit_key,
                )
            )

        total_to_process = (
            len(
                work
            )
        )

        if (
            total_to_process
            !=
            len(
                valid_feedback
            )
        ):
            raise RuntimeError(
                "Not all valid feedback files were added to the processing queue."
            )

        # ====================================================
        # LIVE PROGRESS
        # ====================================================

        st.subheader(
            "Analysis Progress"
        )

        progress = (
            st.progress(
                0.0
            )
        )

        progress_text = (
            st.empty()
        )

        progress_details = (
            st.empty()
        )

        status_box = (
            st.empty()
        )

        progress_text.markdown(
            (
                f"**Starting analysis:** 0 of {total_to_process} "
                f"files processed (0%)."
            )
        )

        progress_details.info(
            (
                f"Remaining: {total_to_process} | "
                f"Successful: 0 | Failed: 0"
            )
        )

        processing_started = (
            time.perf_counter()
        )

        with ThreadPoolExecutor(
            max_workers=MAX_WORKERS
        ) as executor:

            future_map = {
                executor.submit(
                    process_one_feedback,
                    filename,
                    raw_text,
                    memo_sections,
                    specification_text,
                    api_key=api_key,
                    model=DEFAULT_MODEL,
                    classification_prompt=prompt,
                    batch_size=BATCH_SIZE,
                ):
                (
                    filename,
                    result_path,
                    audit_key,
                )

                for (
                    filename,
                    raw_text,
                    result_path,
                    audit_key,
                )
                in work
            }

            completed = 0

            for future in (
                as_completed(
                    future_map
                )
            ):
                (
                    filename,
                    result_path,
                    audit_key,
                ) = (
                    future_map[
                        future
                    ]
                )

                completed += 1

                try:
                    result = (
                        future.result()
                    )

                    stored = {
                        **result,
                        "_audit_key":
                            audit_key,
                    }

                    # Overwrite any previous result for this student.
                    result_path.write_text(
                        json.dumps(
                            stored,
                            indent=2,
                            ensure_ascii=False,
                        ),
                        encoding="utf-8",
                    )

                    results.append(
                        result
                    )

                    status_box.success(
                        f"Completed: {filename}"
                    )

                except Exception as exc:
                    failures.append(
                        (
                            f"{filename}: "
                            f"{type(exc).__name__}: "
                            f"{exc}"
                        )
                    )

                    status_box.error(
                        f"Failed: {filename}"
                    )

                progress_value = (
                    completed
                    /
                    total_to_process
                )

                progress.progress(
                    progress_value
                )

                percent_complete = (
                    progress_value
                    *
                    100
                )

                elapsed_progress = (
                    time.perf_counter()
                    -
                    processing_started
                )

                remaining_files = (
                    total_to_process
                    -
                    completed
                )

                successful_count = (
                    len(
                        results
                    )
                )

                failed_count = (
                    len(
                        failures
                    )
                )

                if completed > 0:
                    average_seconds_per_completed = (
                        elapsed_progress
                        /
                        completed
                    )

                    estimated_remaining_seconds = (
                        average_seconds_per_completed
                        *
                        remaining_files
                    )
                else:
                    estimated_remaining_seconds = 0.0

                progress_text.markdown(
                    (
                        f"**Analysis progress:** "
                        f"{completed} of {total_to_process} files "
                        f"processed ({percent_complete:.0f}%)."
                    )
                )

                progress_details.info(
                    (
                        f"Successful: {successful_count} | "
                        f"Failed: {failed_count} | "
                        f"Remaining: {remaining_files} | "
                        f"Elapsed: {elapsed_progress / 60:.1f} min | "
                        f"Estimated remaining: "
                        f"{estimated_remaining_seconds / 60:.1f} min"
                    )
                )

        progress.progress(
            1.0
        )

        progress_text.markdown(
            "**Analysis progress:** 100% complete."
        )

        progress_details.success(
            (
                f"Analysis finished. "
                f"Successful: {len(results)} | "
                f"Failed: {len(failures)} | "
                f"Total attempted: {total_to_process}"
            )
        )

        elapsed = (
            time.perf_counter()
            -
            start_time
        )

        df = (
            results_to_dataframe(
                results
            )
        )

        run_input_tokens = sum(
            int(
                result.get(
                    "analysis_metadata",
                    {},
                ).get(
                    "input_tokens",
                    0,
                )
            )
            for result
            in results
        )

        run_output_tokens = sum(
            int(
                result.get(
                    "analysis_metadata",
                    {},
                ).get(
                    "output_tokens",
                    0,
                )
            )
            for result
            in results
        )

        run_total_tokens = (
            run_input_tokens
            +
            run_output_tokens
        )

        run_api_requests = sum(
            int(
                result.get(
                    "analysis_metadata",
                    {},
                ).get(
                    "api_requests",
                    0,
                )
            )
            for result
            in results
        )

        actual_run_cost = (
            calculate_model_cost(
                DEFAULT_MODEL,
                run_input_tokens,
                run_output_tokens,
            )
        )

        throughput = (
            len(
                results
            )
            /
            elapsed
            *
            60

            if (
                elapsed > 0
                and
                len(
                    results
                )
                >
                0
            )

            else
            0.0
        )

        meta = {
            "assignment_name":
                assignment_name_clean,

            "assignment_slug":
                slugify(
                    assignment_name_clean
                ),

            "lecturer_user_id":
                st.session_state.get(
                    "user_id"
                ),

            "model_label":
                DEFAULT_MODEL_LABEL,

            "model_id":
                DEFAULT_MODEL,

            "prompt_used":
                prompt,

            "prompt_hash":
                prompt_hash,

            "prompt_modified":
                prompt_modified,

            "workers":
                MAX_WORKERS,

            "batch_size":
                BATCH_SIZE,

            "uploaded_count":
                len(
                    feedback_files
                ),

            "valid_count":
                len(
                    valid_feedback
                ),

            "processed_count":
                len(
                    results
                ),

            "attempted_count":
                total_to_process,

            "cached_count":
                0,

            "invalid_files":
                invalid_feedback,

            "processing_failures":
                failures,

            "failed_partial_items":
                len(
                    df
                ),

            "run_input_tokens":
                run_input_tokens,

            "run_output_tokens":
                run_output_tokens,

            "run_total_tokens":
                run_total_tokens,

            "run_api_requests":
                run_api_requests,

            "actual_run_cost_usd":
                actual_run_cost,

            "dataset_input_tokens":
                run_input_tokens,

            "dataset_output_tokens":
                run_output_tokens,

            "dataset_total_tokens":
                run_total_tokens,

            "dataset_api_requests":
                run_api_requests,

            "dataset_cost_usd":
                actual_run_cost,

            "estimated_time":
                time_estimate,

            "estimated_cost_usd":
                (
                    estimate[
                        "cost_usd"
                    ]
                    if estimate
                    else
                    None
                ),

            "elapsed_seconds":
                elapsed,

            "elapsed_minutes":
                (
                    elapsed
                    /
                    60
                ),

            "throughput_files_per_minute":
                throughput,

            "updated_at":
                time.strftime(
                    "%Y-%m-%d %H:%M:%S"
                ),
        }

        assignment_meta_path(
            assignment_name_clean
        ).write_text(
            json.dumps(
                meta,
                indent=2,
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )

        report_bytes = (
            assignment_report_docx(
                assignment_name_clean,
                results,
                meta,
            )
        )

        report_path = (
            reports_root()
            /
            (
                slugify(
                    assignment_name_clean
                )
                +
                "_EduCodeInsight_Report.docx"
            )
        )

        report_path.write_bytes(
            report_bytes
        )

        st.success(
            "Analysis complete."
        )

        c1, c2, c3, c4 = (
            st.columns(
                4
            )
        )

        c1.metric(
            "Students Processed",
            len(
                results
            ),
        )

        c2.metric(
            "Failed / Partial Items",
            len(
                df
            ),
        )

        c3.metric(
            "Actual Run Cost",
            f"${actual_run_cost:.4f}",
        )

        c4.metric(
            "Actual Runtime",
            f"{elapsed / 60:.1f} min",
        )

        if failures:
            st.warning(
                f"{len(failures)} file(s) failed during processing."
            )

        st.info(
            "Open the Dashboard page to view the assignment "
            "and overall analysis."
        )


# ============================================================
# ASSIGNMENT DASHBOARD
# ============================================================

def render_assignment_dashboard(
    assignment_name: str,
    assignment_slug: str,
):
    results, meta = (
        load_assignment_by_slug(
            assignment_slug
        )
    )

    if not results:
        st.info(
            "No processed student results were found for this assignment."
        )
        return

    df = (
        results_to_dataframe(
            results
        )
    )

    counts = (
        category_counts(
            df
        )
    )

    patterns = (
        rubric_criterion_counts(
            df
        )
    )

    st.subheader(
        assignment_name
    )

    c1, c2, c3, c4 = st.columns(
        4
    )

    c1.metric(
        "Students",
        len(
            results
        ),
    )

    c2.metric(
        "Failed / Partial Items",
        len(
            df
        ),
    )

    c3.metric(
        "AI Cost",
        f"${meta.get('dataset_cost_usd', 0.0):.4f}",
    )

    c4.metric(
        "Processing Time",
        f"{meta.get('elapsed_minutes', 0.0):.1f} min",
    )

    tabs = st.tabs(
        [
            "Overview",
            "Recurring Patterns",
            "Question Analysis",
            "Student Details",
            "AI Usage",
        ]
    )

    with tabs[
        0
    ]:
        st.markdown(
            "#### Error Category Distribution"
        )

        category_df = (
            counts
            .rename(
                "count"
            )
            .to_frame()
        )

        st.bar_chart(
            category_df,
            use_container_width=True,
        )

        st.dataframe(
            category_df.reset_index(
                names="category"
            ),
            use_container_width=True,
            hide_index=True,
        )

    with tabs[
        1
    ]:
        st.markdown(
            "#### Most Frequent Failed / Partial Criteria"
        )

        if patterns.empty:
            st.info(
                "No recurring patterns were found."
            )
        else:
            top_patterns = (
                patterns
                .head(
                    20
                )
            )

            st.bar_chart(
                top_patterns.set_index(
                    "rubric_text"
                )[
                    [
                        "count",
                    ]
                ],
                use_container_width=True,
            )

            st.dataframe(
                top_patterns,
                use_container_width=True,
                hide_index=True,
            )

    with tabs[
        2
    ]:
        st.markdown(
            "#### Question Analysis"
        )

        if df.empty:
            st.info(
                "No failed or partially met criteria were found."
            )
        else:
            question_summary = (
                df
                .groupby(
                    [
                        "question_number",
                        "question_title",
                        "category",
                    ],
                    dropna=False,
                )
                .size()
                .reset_index(
                    name="count"
                )
            )

            pivot = (
                question_summary
                .pivot_table(
                    index=[
                        "question_number",
                        "question_title",
                    ],
                    columns="category",
                    values="count",
                    fill_value=0,
                )
                .reset_index()
            )

            pivot.columns.name = None

            pivot[
                "question_label"
            ] = (
                pivot[
                    "question_number"
                ].astype(
                    str
                )
                +
                " - "
                +
                pivot[
                    "question_title"
                ].fillna(
                    ""
                ).astype(
                    str
                )
            )

            chart_columns = [
                category
                for category
                in CATEGORIES
                if category
                in pivot.columns
            ]

            if chart_columns:

                question_chart = (
                    pivot[
                        [
                            "question_label",
                            *chart_columns,
                        ]
                    ]
                )

                st.bar_chart(
                    question_chart,
                    x="question_label",
                    y=chart_columns,
                    use_container_width=True,
                )

            else:

                st.info(
                    "No error categories were available for the question chart."
                )

            st.dataframe(
                question_summary,
                use_container_width=True,
                hide_index=True,
            )

    with tabs[
        3
    ]:
        st.markdown(
            "#### Student Details"
        )

        student_ids = sorted(
            {
                str(
                    result.get(
                        "student_id",
                        ""
                    )
                )
                for result
                in results
            }
        )

        selected_student = st.selectbox(
            "Student",
            student_ids,
            key=(
                "assignment_student_"
                +
                assignment_slug
            ),
        )

        student_df = (
            df[
                df[
                    "student_id"
                ].astype(
                    str
                )
                ==
                selected_student
            ]
            if not df.empty
            else
            pd.DataFrame()
        )

        if student_df.empty:
            st.info(
                "No failed or partially met criteria were found for this student."
            )
        else:
            st.dataframe(
                student_df[
                    [
                        "question_number",
                        "question_title",
                        "rubric_text",
                        "status",
                        "category",
                        "reason",
                        "justification",
                    ]
                ],
                use_container_width=True,
                hide_index=True,
            )

    with tabs[
        4
    ]:
        st.markdown(
            "#### AI Usage and Reproducibility"
        )

        c1, c2, c3, c4 = st.columns(
            4
        )

        c1.metric(
            "Input Tokens",
            f"{meta.get('dataset_input_tokens', 0):,}",
        )

        c2.metric(
            "Output Tokens",
            f"{meta.get('dataset_output_tokens', 0):,}",
        )

        c3.metric(
            "API Requests",
            f"{meta.get('dataset_api_requests', 0):,}",
        )

        c4.metric(
            "Model",
            meta.get(
                "model_label",
                DEFAULT_MODEL_LABEL,
            ),
        )

        st.write(
            "Prompt hash:",
            meta.get(
                "prompt_hash",
                "Not available",
            ),
        )

        st.write(
            "Prompt modified:",
            meta.get(
                "prompt_modified",
                False,
            ),
        )

        with st.expander(
            "View prompt used"
        ):
            st.code(
                meta.get(
                    "prompt_used",
                    "Prompt not stored.",
                ),
                language="text",
            )


# ============================================================
# OVERALL DASHBOARD
# ============================================================

def render_overall_dashboard():
    assignment_runs = (
        load_all_assignment_runs()
    )

    if not assignment_runs:
        st.info(
            "No stored assignments are available for overall analysis."
        )
        return

    overall = (
        build_overall_analysis(
            assignment_runs
        )
    )

    summary = (
        overall[
            "summary"
        ]
    )

    c1, c2, c3, c4 = st.columns(
        4
    )

    c1.metric(
        "Assignments",
        summary[
            "assignments_analysed"
        ],
    )

    c2.metric(
        "Student Submissions",
        summary[
            "student_submissions_processed"
        ],
    )

    c3.metric(
        "Unique Students",
        summary[
            "unique_students"
        ],
    )

    c4.metric(
        "Failed / Partial Items",
        summary[
            "failed_partial_items"
        ],
    )

    c1, c2, c3 = st.columns(
        3
    )

    c1.metric(
        "Total Tokens",
        f"{summary['total_tokens']:,}",
    )

    c2.metric(
        "Total AI Cost",
        f"${summary['total_ai_cost_usd']:.4f}",
    )

    c3.metric(
        "Total Processing Time",
        f"{summary['total_processing_minutes']:.1f} min",
    )

    tabs = st.tabs(
        [
            "Overview",
            "Assignment Comparison",
            "Category Trends",
            "Recurring Errors",
            "Student Trends",
            "AI Usage & Cost",
        ]
    )

    with tabs[
        0
    ]:
        st.markdown(
            "#### Overall Error Category Distribution"
        )

        category_distribution = (
            overall[
                "category_distribution"
            ]
        )

        if category_distribution.empty:
            st.info(
                "No category data is available."
            )
        else:
            st.bar_chart(
                category_distribution
                .set_index(
                    "category"
                )[
                    [
                        "count",
                    ]
                ],
                use_container_width=True,
            )

            st.dataframe(
                category_distribution,
                use_container_width=True,
                hide_index=True,
            )

    with tabs[
        1
    ]:
        st.markdown(
            "#### Assignment Comparison"
        )

        assignment_summary = (
            overall[
                "assignment_summary"
            ]
        )

        if not assignment_summary.empty:
            st.bar_chart(
                assignment_summary
                .set_index(
                    "assignment"
                )[
                    [
                        "failed_partial_items",
                    ]
                ],
                use_container_width=True,
            )

            st.dataframe(
                assignment_summary,
                use_container_width=True,
                hide_index=True,
            )

    with tabs[
        2
    ]:
        st.markdown(
            "#### Error Category Trends by Assignment"
        )

        category_by_assignment = (
            overall[
                "category_by_assignment"
            ]
        )

        if category_by_assignment.empty:
            st.info(
                "No category trend data is available."
            )
        else:
            pivot = (
                category_by_assignment
                .pivot_table(
                    index="assignment",
                    columns="category",
                    values="count",
                    fill_value=0,
                )
                .reset_index()
            )

            pivot.columns.name = None

            chart_columns = [
                category
                for category
                in CATEGORIES
                if category
                in pivot.columns
            ]

            if chart_columns:

                st.bar_chart(
                    pivot,
                    x="assignment",
                    y=chart_columns,
                    use_container_width=True,
                )

            else:

                st.info(
                    "No error categories were available for the assignment trend chart."
                )

            st.dataframe(
                category_by_assignment,
                use_container_width=True,
                hide_index=True,
            )

    with tabs[
        3
    ]:
        st.markdown(
            "#### Recurring Errors Across Assignments"
        )

        recurring_errors = (
            overall[
                "recurring_errors"
            ]
        )

        if recurring_errors.empty:
            st.info(
                "No recurring error data is available."
            )
        else:
            top = (
                recurring_errors
                .head(
                    20
                )
            )

            st.bar_chart(
                top.set_index(
                    "rubric_text"
                )[
                    [
                        "students_affected",
                    ]
                ],
                use_container_width=True,
            )

            st.dataframe(
                top,
                use_container_width=True,
                hide_index=True,
            )

    with tabs[
        4
    ]:
        st.markdown(
            "#### Student Trends Across Assignments"
        )

        student_overall = (
            overall[
                "student_overall"
            ]
        )

        if student_overall.empty:
            st.info(
                "No student trend data is available."
            )
        else:
            st.dataframe(
                student_overall,
                use_container_width=True,
                hide_index=True,
            )

    with tabs[
        5
    ]:
        st.markdown(
            "#### AI Usage and Cost"
        )

        model_usage = (
            overall[
                "model_usage"
            ]
        )

        assignment_summary = (
            overall[
                "assignment_summary"
            ]
        )

        if not model_usage.empty:
            st.dataframe(
                model_usage,
                use_container_width=True,
                hide_index=True,
            )

        if (
            not assignment_summary.empty
            and
            "ai_cost_usd"
            in assignment_summary.columns
        ):
            st.bar_chart(
                assignment_summary
                .set_index(
                    "assignment"
                )[
                    [
                        "ai_cost_usd",
                    ]
                ],
                use_container_width=True,
            )


# ============================================================
# DASHBOARD PAGE
# ============================================================

def render_dashboard():
    st.header(
        "Dashboard"
    )

    assignment_view, overall_view = (
        st.tabs(
            [
                "Assignment Analysis",
                "Overall Analysis",
            ]
        )
    )

    with assignment_view:
        assignments = (
            list_stored_assignments()
        )

        if not assignments:
            st.info(
                "No processed assignments are available yet."
            )
        else:
            label_to_slug = {
                name:
                    slug
                for (
                    name,
                    slug,
                )
                in assignments
            }

            selected_assignment = (
                st.selectbox(
                    "Assignment",
                    list(
                        label_to_slug.keys()
                    ),
                )
            )

            render_assignment_dashboard(
                selected_assignment,
                label_to_slug[
                    selected_assignment
                ],
            )

    with overall_view:
        render_overall_dashboard()


# ============================================================
# REPORTS PAGE
# ============================================================

def render_reports():
    st.header(
        "Reports"
    )

    st.write(
        "Reports are generated as concise Microsoft Word documents "
        "with visual summaries and the main recurring areas of difficulty."
    )

    assignments = (
        list_stored_assignments()
    )

    if not assignments:
        st.info(
            "No reports are available yet."
        )
        return

    label_to_slug = {
        name:
            slug
        for (
            name,
            slug,
        )
        in assignments
    }

    selected_assignment = (
        st.selectbox(
            "Assignment report",
            list(
                label_to_slug.keys()
            ),
            key="report_assignment",
        )
    )

    results, meta = (
        load_assignment_by_slug(
            label_to_slug[
                selected_assignment
            ]
        )
    )

    assignment_report = (
        assignment_report_docx(
            selected_assignment,
            results,
            meta,
        )
    )

    st.download_button(
        "Download Assignment Report",
        data=assignment_report,
        file_name=(
            slugify(
                selected_assignment
            )
            +
            "_EduCodeInsight_Report.docx"
        ),
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        use_container_width=True,
    )


    overall_runs = (
        load_all_assignment_runs()
    )

    if overall_runs:

        overall = (
            build_overall_analysis(
                overall_runs
            )
        )

        overall_report = (
            overall_report_docx(
                overall
            )
        )

        st.download_button(
            "Download Overall Analysis Report",
            data=overall_report,
            file_name="EduCodeInsight_Overall_Report.docx",
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            use_container_width=True,
        )


