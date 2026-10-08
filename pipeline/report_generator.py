"""DOCX report generation service."""
from __future__ import annotations
import io, re
from pathlib import Path
from typing import Optional
import pandas as pd
import matplotlib.pyplot as plt
from .models import CATEGORIES
from .config import DEFAULT_MODEL_LABEL, DEFAULT_MODEL
from .analysis_engine import AnalysisEngine
from .chart_generator import ChartGenerator

class ReportGenerator:
    """Generate assignment and cross-assignment reports independently of the GUI."""
    def __init__(self, output_folder: Path, analyzer: AnalysisEngine | None = None,
                 charts: ChartGenerator | None = None):
        self.output_folder = Path(output_folder); self.output_folder.mkdir(parents=True, exist_ok=True)
        self.analyzer = analyzer or AnalysisEngine(); self.charts = charts or ChartGenerator()
        try:
            from docx import Document
            from docx.shared import Inches
        except ImportError as exc:
            raise RuntimeError("python-docx is required for DOCX reports. Install it with: pip install python-docx") from exc
        self.Document = Document; self.Inches = Inches

    def _add_title(self, doc, title: str): doc.add_heading(title, level=1)

    def _add_key_value_table(self, doc, rows):
        table = doc.add_table(rows=0, cols=2); table.style = "Light Grid Accent 1"
        for key, value in rows:
            cells = table.add_row().cells; cells[0].text = str(key); cells[1].text = str(value)
        return table

    def _add_figure(self, doc, fig, width=6.4):
        buffer = io.BytesIO(); fig.savefig(buffer, format="png", dpi=160, bbox_inches="tight")
        plt.close(fig); buffer.seek(0); doc.add_picture(buffer, width=self.Inches(width))

    def _add_category_table(self, doc, counts):
        total = int(counts.sum()) if hasattr(counts, "sum") else 0
        table = doc.add_table(rows=1, cols=3); table.style = "Light Grid Accent 1"
        for i, h in enumerate(["Category", "Count", "Percentage"]): table.rows[0].cells[i].text = h
        for category in CATEGORIES:
            count = int(counts.get(category, 0)); pct = (count / total * 100) if total else 0
            cells = table.add_row().cells; cells[0].text = category; cells[1].text = f"{count:,}"; cells[2].text = f"{pct:.1f}%"

    def generate_assignment_report(self, run_result: dict, assignment_name: str = "Assignment",
                                   filename: Optional[str] = None) -> Path:
        analysis = self.analyzer.assignment_analysis(run_result)
        df, counts, patterns = analysis["dataframe"], analysis["category_counts"], analysis["recurring_errors"]
        doc = self.Document(); self._add_title(doc, f"EduCodeInsight - {assignment_name} Report")
        total_items = len(df); students = len(run_result.get("results", []))
        dominant = counts.idxmax() if total_items and counts.sum() else "None"
        dominant_count = int(counts.max()) if total_items and counts.sum() else 0
        doc.add_heading("1. Executive Summary", level=2)
        doc.add_paragraph(f"EduCodeInsight analysed {students} student submissions and identified {total_items:,} failed or partially met rubric criteria. The dominant error category was {dominant} ({dominant_count:,} items).")
        doc.add_heading("2. Visual Overview", level=2)
        self._add_figure(doc, self.charts.category_distribution(counts)); self._add_figure(doc, self.charts.errors_by_question(df)); self._add_figure(doc, self.charts.top_recurring_errors(patterns))
        doc.add_heading("3. What Students Mainly Got Wrong", level=2)
        top = patterns.head(10)
        if top.empty: doc.add_paragraph("No recurring failed or partially met criteria were identified.")
        else:
            table = doc.add_table(rows=1, cols=3); table.style = "Light Grid Accent 1"
            for i, h in enumerate(["Problem", "Students affected", "Dominant type"]): table.rows[0].cells[i].text = h
            for _, row in top.iterrows():
                cells = table.add_row().cells; cells[0].text = str(row["rubric_text"]); cells[1].text = str(int(row["count"])); cells[2].text = str(row["dominant_category"])
        doc.add_heading("4. Error Category Summary", level=2); self._add_category_table(doc, counts)
        doc.add_heading("5. Analysis Details", level=2)
        self._add_key_value_table(doc, [
            ("Model", run_result.get("model_label", DEFAULT_MODEL_LABEL)), ("Model ID", run_result.get("model_id", DEFAULT_MODEL)),
            ("Prompt modified", run_result.get("prompt_modified", False)), ("Prompt hash", run_result.get("prompt_hash", "")),
            ("Valid files", run_result.get("valid_count", 0)), ("Processing failures", len(run_result.get("failures", []))),
            ("Processing time", f"{run_result.get('elapsed_minutes', 0.0):.1f} minutes"),
            ("Throughput", f"{run_result.get('throughput_files_per_minute', 0.0):.1f} files/min"),
            ("Input tokens", f"{run_result.get('run_input_tokens', 0):,}"), ("Output tokens", f"{run_result.get('run_output_tokens', 0):,}"),
            ("Total tokens", f"{run_result.get('run_total_tokens', 0):,}"), ("API requests", f"{run_result.get('run_api_requests', 0):,}"),
            ("Actual run cost", f"${run_result.get('actual_run_cost_usd', 0.0):.6f}"),
        ])
        path = self.output_folder / (filename or f"{re.sub(r'[^A-Za-z0-9_-]+','_',assignment_name)}_EduCodeInsight_Report.docx")
        doc.save(path); return path

    def generate_overall_report(self, overall: dict, filename: str = "EduCodeInsight_Overall_Report.docx") -> Path:
        doc = self.Document(); self._add_title(doc, "EduCodeInsight - Overall Analysis Report")
        summary = overall.get("summary", {})
        doc.add_heading("1. Executive Summary", level=2)
        doc.add_paragraph(f"The overall analysis combines {summary.get('assignments_analysed', 0)} assignments, {summary.get('student_submissions_processed', 0)} student submissions, and {summary.get('failed_partial_items', 0):,} failed or partially met criteria.")
        doc.add_heading("2. Visual Overview", level=2)
        distribution = overall.get("category_distribution", pd.Series(dtype=int)); self._add_figure(doc, self.charts.category_distribution(distribution))
        assignment_summary = overall.get("assignment_summary", pd.DataFrame())
        if not assignment_summary.empty and "failed_partial_items" in assignment_summary.columns:
            fig, ax = plt.subplots(figsize=(10, 5)); ax.bar(assignment_summary["assignment"].astype(str), assignment_summary["failed_partial_items"].values)
            ax.set_ylabel("Failed / partially met items"); ax.set_title("Failed / Partially Met Items by Assignment"); ax.tick_params(axis="x", rotation=30); fig.tight_layout(); self._add_figure(doc, fig)
        doc.add_heading("3. Main Recurring Problems", level=2); recurring = overall.get("recurring_errors", pd.DataFrame())
        if not recurring.empty:
            table = doc.add_table(rows=1, cols=3); table.style = "Light Grid Accent 1"
            for i, h in enumerate(["Problem", "Occurrences", "Dominant type"]): table.rows[0].cells[i].text = h
            for _, row in recurring.head(10).iterrows():
                cells = table.add_row().cells; cells[0].text = str(row.get("rubric_text", "")); cells[1].text = str(int(row.get("count", 0))); cells[2].text = str(row.get("dominant_category", ""))
        else: doc.add_paragraph("No recurring problems were identified.")
        doc.add_heading("4. Assignment Comparison", level=2)
        if not assignment_summary.empty: self._add_key_value_table(doc, [(row.get("assignment", ""), f"{int(row.get('failed_partial_items', 0)):,} failed/partial items") for _, row in assignment_summary.iterrows()])
        doc.add_heading("5. Analysis Details", level=2)
        self._add_key_value_table(doc, [("Total input tokens", f"{summary.get('input_tokens', 0):,}"), ("Total output tokens", f"{summary.get('output_tokens', 0):,}"), ("Total tokens", f"{summary.get('total_tokens', 0):,}"), ("Total API requests", f"{summary.get('api_requests', 0):,}"), ("Total AI cost", f"${summary.get('total_ai_cost_usd', 0.0):.6f}"), ("Total processing time", f"{summary.get('total_processing_minutes', 0.0):.1f} minutes")])
        path = self.output_folder / filename; doc.save(path); return path

__all__ = ["ReportGenerator"]
